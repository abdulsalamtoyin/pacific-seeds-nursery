"""Turn a grid from the app into a real .xlsx file.

The desktop app's tables are HTML, but the client works in Excel: packet lists
feed Bartender, and the replacement log has a two-band header that has to come
out looking like `Tab information.xlsx`. Rather than bundle a spreadsheet
writer into the browser, the grid posts what it is showing and this builds the
workbook with openpyxl, which the backend already depends on.

Kept separate from app.py so the layout rules can be tested without HTTP.
"""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Excel forbids these in a sheet name, caps it at 31 characters, and will not
# accept an empty one. "History" is reserved.
_ILLEGAL = re.compile(r"[\[\]:*?/\\]")
_MAX_NAME = 31
_RESERVED = {"history"}


def sanitise_sheet_name(name: str, taken: set[str] | None = None) -> str:
    """A sheet name Excel will accept, unique against `taken`.

    Tab names like "Replacements" fit, but a long nursery code appended to
    "Packet Prep 1" would not, and a user-added tab could contain a slash.
    """
    cleaned = _ILLEGAL.sub("-", str(name or "")).strip().strip("'")
    if not cleaned:
        cleaned = "Sheet"
    if cleaned.lower() in _RESERVED:
        cleaned = f"{cleaned}_"
    cleaned = cleaned[:_MAX_NAME]

    if taken is None:
        return cleaned

    # Uniquify by suffixing " (2)", trimming the stem so the result still fits.
    candidate = cleaned
    n = 2
    lowered = {t.lower() for t in taken}
    while candidate.lower() in lowered:
        suffix = f" ({n})"
        candidate = cleaned[:_MAX_NAME - len(suffix)] + suffix
        n += 1
    return candidate


def _rgb(colour: Any) -> str | None:
    """'#ffd9d9' or 'ffd9d9' -> 'FFFFD9D9' (openpyxl wants aRGB)."""
    if not colour:
        return None
    text = str(colour).lstrip("#").strip()
    if len(text) == 3:                       # #abc -> aabbcc
        text = "".join(c * 2 for c in text)
    if len(text) == 6:
        text = f"FF{text}"
    if len(text) != 8 or not re.fullmatch(r"[0-9A-Fa-f]{8}", text):
        return None
    return text.upper()


def _cell_ref(key: Any) -> tuple[int, int] | None:
    """'2,3' -> (2, 3); anything malformed is ignored rather than fatal."""
    try:
        row, col = (int(part) for part in str(key).split(","))
    except ValueError:
        return None
    if row < 0 or col < 0:
        return None
    return row, col


def _write_sheet(ws, sheet: dict) -> None:
    headers: list[str] = list(sheet.get("headers") or [])
    rows: list[list[Any]] = list(sheet.get("rows") or [])
    bands = sheet.get("bands") or []
    colours = sheet.get("colours") or {}
    styles = sheet.get("styles") or {}
    widths = sheet.get("widths") or []

    header_row = 1
    if bands:
        # A band spans several columns above the real headers, as in the
        # client's format: Original entry | Replaced entry | Reason.
        col = 1
        for label, span in bands:
            span = max(1, int(span))
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            if span > 1:
                ws.merge_cells(start_row=1, start_column=col,
                               end_row=1, end_column=col + span - 1)
            col += span
        header_row = 2

    for i, head in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=head).font = Font(bold=True)

    first_data_row = header_row + 1
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.cell(row=first_data_row + r, column=c + 1, value=value)

    # Fills and text styles are addressed by the position the user saw them in,
    # zero-based over the data rows only.
    for key, colour in colours.items():
        ref = _cell_ref(key)
        rgb = _rgb(colour)
        if not ref or not rgb:
            continue
        r, c = ref
        ws.cell(row=first_data_row + r, column=c + 1).fill = PatternFill(
            start_color=rgb, end_color=rgb, fill_type="solid")

    for key, style in styles.items():
        ref = _cell_ref(key)
        if not ref:
            continue
        r, c = ref
        cell = ws.cell(row=first_data_row + r, column=c + 1)
        size = style.get("size")
        cell.font = Font(
            bold=bool(style.get("bold")),
            italic=bool(style.get("italic")),
            size=float(size) if size else None,
            color=_rgb(style.get("colour")),
        )
        if style.get("align") in {"left", "center", "right"}:
            cell.alignment = Alignment(horizontal=style["align"])

    for i, width in enumerate(widths, start=1):
        try:
            ws.column_dimensions[get_column_letter(i)].width = float(width)
        except (TypeError, ValueError):
            continue

    # Excel's own filter dropdowns, so the exported sheet behaves like the tab.
    if headers and rows:
        last = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A{header_row}:{last}{header_row + len(rows)}"
    ws.freeze_panes = ws.cell(row=first_data_row, column=1)


def build_workbook(payload: dict) -> Workbook:
    """Build a workbook from `{filename, sheets: [...]}`."""
    sheets = payload.get("sheets") or []
    if not sheets:
        raise ValueError("Nothing to export — no sheets were supplied.")

    # An export that would produce a file of nothing but headers is almost
    # always a mistake, so say so rather than hand over an empty spreadsheet.
    if all(not (s.get("rows") or []) for s in sheets):
        names = ", ".join(str(s.get("name", "?")) for s in sheets)
        raise ValueError(f"Nothing to export — {names} has no rows yet.")

    wb = Workbook()
    wb.remove(wb.active)
    taken: set[str] = set()
    for sheet in sheets:
        name = sanitise_sheet_name(sheet.get("name"), taken)
        taken.add(name)
        _write_sheet(wb.create_sheet(title=name), sheet)
    return wb


def workbook_bytes(payload: dict) -> bytes:
    buffer = BytesIO()
    build_workbook(payload).save(buffer)
    return buffer.getvalue()
