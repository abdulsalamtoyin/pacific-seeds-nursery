"""The exported .xlsx has to match the format the client sent us.

`Tab information.xlsx` shows Replacements as three bands — Original entry,
Replaced entry, Reason — merged above the real headers. Packet lists feed
Bartender, so their column order and row order have to be exact. These build
the workbook in memory and read it back rather than trusting the writer.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.xlsx_export import (
    build_workbook, sanitise_sheet_name, workbook_bytes,
)

# The nine-column entry block from Tab information.xlsx, in order.
ENTRY_BLOCK = [
    "QR Key", "Source ID", "Material ID", "Inbred Code",
    "Experimental Hybrid Code", "Pedigree", "Generation", "CMS reaction",
    "W/column",
]


def roundtrip(payload: dict):
    """Build the workbook and read it back as openpyxl would see it."""
    return load_workbook(BytesIO(workbook_bytes(payload)))


def simple(name="Sheet1", rows=None, **extra) -> dict:
    sheet = {
        "name": name,
        "headers": ["Range", "Row", "Source ID"],
        "rows": rows if rows is not None else [[1, 2, "SRC-0001"]],
    }
    sheet.update(extra)
    return {"filename": "t", "sheets": [sheet]}


# ------------------------------------------------------------------ basics


def test_headers_and_rows_land_where_expected():
    ws = roundtrip(simple())["Sheet1"]
    assert [c.value for c in ws[1]] == ["Range", "Row", "Source ID"]
    assert [c.value for c in ws[2]] == [1, 2, "SRC-0001"]


def test_headers_are_bold():
    ws = roundtrip(simple())["Sheet1"]
    assert ws["A1"].font.bold is True


def test_multiple_sheets_become_multiple_tabs():
    wb = roundtrip({"filename": "book", "sheets": [
        {"name": "Fieldbook", "headers": ["Range"], "rows": [[1]]},
        {"name": "Nursery list", "headers": ["Source ID"], "rows": [["S1"]]},
    ]})
    assert wb.sheetnames == ["Fieldbook", "Nursery list"]


def test_row_order_is_preserved_exactly():
    """Bartender reads the packets in the order we write them."""
    rows = [[1, 1, "SRC-0003"], [1, 2, "SRC-0002"], [1, 3, "SRC-0001"]]
    ws = roundtrip(simple(rows=rows))["Sheet1"]
    assert [ws.cell(row=r, column=3).value for r in (2, 3, 4)] == [
        "SRC-0003", "SRC-0002", "SRC-0001"]


# ------------------------------------------------------------------ bands


def test_bands_merge_above_the_headers():
    """Original entry | Replaced entry | Reason, as in Tab information.xlsx."""
    payload = {"filename": "r", "sheets": [{
        "name": "Replacements",
        "bands": [["Original entry", 9], ["Replaced entry", 9], ["Reason", 1]],
        "headers": ENTRY_BLOCK + ENTRY_BLOCK + ["Reason"],
        "rows": [["1_2"] + [""] * 8 + ["1_3"] + [""] * 8 + ["mouse damage"]],
    }]}
    ws = roundtrip(payload)["Replacements"]

    assert ws["A1"].value == "Original entry"
    assert ws["J1"].value == "Replaced entry"
    assert ws["S1"].value == "Reason"
    # Real headers sit on row 2, data from row 3 — matching the client's sheet.
    assert ws["A2"].value == "QR Key"
    assert ws["B2"].value == "Source ID"
    assert ws["A3"].value == "1_2"

    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:I1" in merged
    assert "J1:R1" in merged


def test_a_single_column_band_is_not_merged():
    ws = roundtrip({"filename": "e", "sheets": [{
        "name": "Planting errors",
        "bands": [["Original entry", 9], ["Reason", 1]],
        "headers": ENTRY_BLOCK + ["Reason"],
        "rows": [["1_2"] + [""] * 8 + ["wrong plot"]],
    }]})["Planting errors"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:I1" in merged
    assert not any(r.startswith("J1:J1") for r in merged)


def test_without_bands_headers_stay_on_row_one():
    ws = roundtrip(simple())["Sheet1"]
    assert ws["A1"].value == "Range"


# ------------------------------------------------------------------ styling


def test_cell_fill_survives_the_round_trip():
    ws = roundtrip(simple(colours={"0,2": "#ffd9d9"}))["Sheet1"]
    # Row 0, column 2 of the data == C2 in the sheet.
    assert ws["C2"].fill.start_color.rgb == "FFFFD9D9"


def test_short_hex_and_bare_hex_are_both_accepted():
    ws = roundtrip(simple(colours={"0,0": "#abc", "0,1": "ffd9d9"}))["Sheet1"]
    assert ws["A2"].fill.start_color.rgb == "FFAABBCC"
    assert ws["B2"].fill.start_color.rgb == "FFFFD9D9"


def test_a_nonsense_colour_is_ignored_not_fatal():
    ws = roundtrip(simple(colours={"0,0": "rebeccapurple", "0,1": ""}))["Sheet1"]
    assert ws["A2"].value == 1


def test_bold_and_italic_survive():
    ws = roundtrip(simple(styles={"0,0": {"bold": True, "italic": True}}))["Sheet1"]
    assert ws["A2"].font.bold is True
    assert ws["A2"].font.italic is True


def test_malformed_cell_references_are_skipped():
    ws = roundtrip(simple(colours={"nope": "#ffffff", "-1,0": "#ffffff"}))["Sheet1"]
    assert ws["A2"].value == 1


def test_column_widths_are_applied():
    ws = roundtrip(simple(widths=[12.5, 6, 28.75]))["Sheet1"]
    assert ws.column_dimensions["A"].width == pytest.approx(12.5)
    assert ws.column_dimensions["C"].width == pytest.approx(28.75)


def test_autofilter_covers_the_header_and_data():
    ws = roundtrip(simple(rows=[[1, 1, "a"], [1, 2, "b"]]))["Sheet1"]
    assert ws.auto_filter.ref == "A1:C3"


def test_autofilter_starts_below_a_band_row():
    ws = roundtrip({"filename": "r", "sheets": [{
        "name": "R",
        "bands": [["Original entry", 3]],
        "headers": ["QR Key", "Source ID", "Reason"],
        "rows": [["1_2", "S1", "x"]],
    }]})["R"]
    assert ws.auto_filter.ref == "A2:C3"


# ------------------------------------------------------------------ names


@pytest.mark.parametrize("raw,expected", [
    ("Fieldbook", "Fieldbook"),
    ("Packet Prep 1", "Packet Prep 1"),
    ("Replacements/Errors", "Replacements-Errors"),
    ("a[b]c:d*e?f", "a-b-c-d-e-f"),
    ("", "Sheet"),
    ("History", "History_"),
])
def test_sheet_names_are_made_legal(raw, expected):
    assert sanitise_sheet_name(raw) == expected


def test_long_sheet_names_are_truncated_to_31():
    name = sanitise_sheet_name("Packet Prep for " + "X" * 40)
    assert len(name) == 31


def test_duplicate_sheet_names_are_uniquified():
    wb = roundtrip({"filename": "b", "sheets": [
        {"name": "Packet Prep", "headers": ["a"], "rows": [[1]]},
        {"name": "Packet Prep", "headers": ["a"], "rows": [[2]]},
    ]})
    assert wb.sheetnames == ["Packet Prep", "Packet Prep (2)"]


def test_uniquified_long_names_still_fit():
    taken = {"X" * 31}
    name = sanitise_sheet_name("X" * 40, taken)
    assert len(name) <= 31
    assert name not in taken


# ------------------------------------------------------------------ refusals


def test_an_export_with_no_sheets_is_refused():
    with pytest.raises(ValueError, match="no sheets"):
        build_workbook({"filename": "x", "sheets": []})


def test_an_export_where_every_sheet_is_empty_is_refused():
    """Better to say why than to hand over a file of nothing but headers."""
    with pytest.raises(ValueError, match="no rows yet"):
        build_workbook({"filename": "x", "sheets": [
            {"name": "Fieldbook", "headers": ["Range"], "rows": []},
        ]})


def test_one_empty_sheet_alongside_a_full_one_is_allowed():
    """A whole-book export should not fail because one tab is untouched."""
    wb = roundtrip({"filename": "b", "sheets": [
        {"name": "Fieldbook", "headers": ["Range"], "rows": [[1]]},
        {"name": "Comments", "headers": ["Stage"], "rows": []},
    ]})
    assert wb.sheetnames == ["Fieldbook", "Comments"]


# ------------------------------------------------------------------ values


def test_dates_are_written_as_the_iso_text_they_arrived_as():
    """The app stores ISO strings; coercing them to Excel serials would
    silently reformat what the user typed."""
    ws = roundtrip(simple(rows=[[1, 2, "2026-08-20"]]))["Sheet1"]
    assert ws["C2"].value == "2026-08-20"


def test_blank_cells_stay_blank():
    ws = roundtrip(simple(rows=[[1, 2, ""]]))["Sheet1"]
    assert ws["C2"].value in (None, "")
