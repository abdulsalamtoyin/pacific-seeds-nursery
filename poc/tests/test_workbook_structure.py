"""The generated workbook must match nursery_spec.json exactly.

Layout assertions for Nursery data come from the client's File for Toyin.xlsx.
"""
import openpyxl
import pytest

from excel_workflow.build_workbooks import build_nursery_template
from excel_workflow.spec.loader import load_spec, tabs_for


@pytest.fixture(scope="module")
def selection_book():
    path = build_nursery_template(nursery_types=["Selection"], planting_dates=2)
    return openpyxl.load_workbook(path)


@pytest.fixture(scope="module")
def ab_book():
    path = build_nursery_template(nursery_types=["AB"], planting_dates=1)
    return openpyxl.load_workbook(path)


# --- Task 6: tab structure ----------------------------------------------------

def test_sheet_names_match_the_spec_exactly(selection_book):
    assert selection_book.sheetnames == tabs_for(["Selection"], planting_dates=2)


def test_settings_sheet_is_hidden(selection_book):
    assert selection_book["Settings"].sheet_state == "hidden"


def test_old_tab_names_are_gone(selection_book):
    for stale in ("Map", "Replacements done", "Planting error noted", "Packet Prep"):
        assert stale not in selection_book.sheetnames


def test_material_map_precedes_field_map(selection_book):
    names = selection_book.sheetnames
    assert names.index("Material Map") == names.index("Field Map") - 1


def test_ab_workbook_includes_date_recording_before_operations(ab_book):
    names = ab_book.sheetnames
    assert names == tabs_for(["AB"], planting_dates=1)
    assert names.index("Date recording") == names.index("Operations") - 1


def test_selection_workbook_has_no_ab_extras(selection_book):
    for extra in ("BC0 labels", "Pulling bags", "Date recording"):
        assert extra not in selection_book.sheetnames


def test_packet_prep_has_split_no_column(selection_book):
    headers = [c.value for c in selection_book["Packet Prep 1"][5] if c.value]
    assert "Split no." in headers


# --- Task 7: Nursery data layout ---------------------------------------------

def test_nursery_data_title_is_merged_and_bold(selection_book):
    ws = selection_book["Nursery data"]
    assert "A1:D1" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["A1"].value == "R&D Fieldbook - Grain Sorghum"
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.size == 20


def test_nursery_data_rows_three_and_four_are_merged_but_blank(selection_book):
    ws = selection_book["Nursery data"]
    merged = [str(r) for r in ws.merged_cells.ranges]
    assert "B3:C3" in merged and "B4:C4" in merged
    assert ws["B3"].value is None
    assert ws["B4"].value is None


def test_nursery_data_labels_are_in_column_b_with_blank_column_c(selection_book):
    ws = selection_book["Nursery data"]
    assert ws["B5"].value == "GPS coordinates"
    assert ws["B23"].value == "Harvest date"
    for row in range(5, 24):
        assert ws.cell(row=row, column=3).value is None


def test_nursery_data_has_all_nineteen_labels(selection_book):
    ws = selection_book["Nursery data"]
    labels = [ws.cell(row=r, column=2).value for r in range(5, 24)]
    assert labels == load_spec()["nursery_data_labels"]


def test_nursery_data_column_widths_match_the_sample(selection_book):
    ws = selection_book["Nursery data"]
    assert round(ws.column_dimensions["B"].width, 2) == 58.29
    assert round(ws.column_dimensions["G"].width, 2) == 59.57


def test_nursery_data_footer_row_is_merged(selection_book):
    ws = selection_book["Nursery data"]
    assert "A24:D24" in [str(r) for r in ws.merged_cells.ranges]


# --- Task 8: Replacements and Errors, Operations, Comments -------------------

def _headers(ws, row):
    return [c.value for c in ws[row] if c.value is not None]


def test_replacements_has_stage_first_and_split_before_status(selection_book):
    headers = _headers(selection_book["Replacements and Errors"], 5)
    assert headers[0] == "Stage"
    assert headers.index("Split no.") == headers.index("Status") - 1


def test_replacements_qr_column_is_wide_enough_for_150_chars(selection_book):
    ws = selection_book["Replacements and Errors"]
    headers = _headers(ws, 5)
    qr_letter = chr(ord("A") + headers.index("QR"))
    assert ws.column_dimensions[qr_letter].width >= 150


def test_replacements_status_defaults_to_waiting_decision(selection_book):
    ws = selection_book["Replacements and Errors"]
    headers = _headers(ws, 5)
    status_letter = chr(ord("A") + headers.index("Status"))
    assert ws[f"{status_letter}6"].value == "Waiting decision"


def test_replacements_has_no_rename_this_tab_row(selection_book):
    ws = selection_book["Replacements and Errors"]
    text = " ".join(str(c.value) for r in ws.iter_rows(max_row=6) for c in r)
    assert "Rename this tab" not in text


def test_operations_uses_group_headings_and_keeps_stage(selection_book):
    ws = selection_book["Operations"]
    assert ws["A1"].value == "Operations Overview"
    assert _headers(ws, 2) == ["Stage", "Group 1", "Group 2", "Group 3"]


def test_operations_dropped_the_old_column_names(selection_book):
    headers = _headers(selection_book["Operations"], 2)
    assert "Reminders" not in headers
    assert "Plan" not in headers


def test_operations_has_no_empty_second_row(selection_book):
    # Title on row 1, headers immediately on row 2 — no gap.
    ws = selection_book["Operations"]
    assert ws["A2"].value == "Stage"


def test_comments_mirrors_operations_with_its_own_title(selection_book):
    ws = selection_book["Comments"]
    assert ws["A1"].value == "Field Comments"
    assert _headers(ws, 2) == ["Stage", "Group 1", "Group 2", "Group 3"]


def test_grouped_sheets_have_all_borders(selection_book):
    ws = selection_book["Operations"]
    for row in range(2, 10):
        for col in range(1, 5):
            border = ws.cell(row=row, column=col).border
            assert border.top.style == "thin", f"row {row} col {col} top"
            assert border.left.style == "thin", f"row {row} col {col} left"
